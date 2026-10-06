
"""
==========================================================
HYBRID COMPOSITE PREDICTION SYSTEM (HCPS)

reports.py

Professional Engineering Report Export Module

Exports:
- Excel Engineering Report
- CSV Data File
- Text Engineering Summary

Current Experimental Properties:
1. Flexural Strength (MPa)
2. Flexural Modulus (GPa)
3. Compressive Strength (MPa)
4. Impact Strength (kJ/m²)

AI Prediction Properties:
1. Composite Density (g/cm³)
2. Flexural Strength (MPa)
3. Flexural Modulus (GPa)
4. Compressive Strength (MPa)
5. Impact Strength (kJ/m²)

Author: Daniel Thomas
==========================================================
"""

import csv
import os

from datetime import datetime
from tkinter import filedialog

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment


# ==========================================================
# CURRENT HCPS EXPERIMENTAL PROPERTIES
# ==========================================================

EXPERIMENTAL_PROPERTIES = [
    "Flexural Strength (MPa)",
    "Flexural Modulus (GPa)",
    "Compressive Strength (MPa)",
    "Impact Strength (kJ/m²)",
]


# ==========================================================
# CURRENT HCPS AI PREDICTION PROPERTIES
# ==========================================================

PREDICTION_PROPERTIES = [
    "Composite Density (g/cm³)",
    "Flexural Strength (MPa)",
    "Flexural Modulus (GPa)",
    "Compressive Strength (MPa)",
    "Impact Strength (kJ/m²)",
]


# ==========================================================
# OLD / INVALID PROPERTIES
#
# These properties must NEVER appear in the HCPS export.
# ==========================================================

INVALID_PROPERTIES = {
    "Young's Modulus (GPa)",
    "Young’s Modulus (GPa)",
    "Hardness (Shore D)",
    "Hardness",
    "Penetration Depth (mm)",
    "Penetration Resistance",
    "Penetration Resistance (index)",
}


# ==========================================================
# FLATTEN HCPS DATA
# ==========================================================

def flatten_results(data):
    """
    Convert the HCPS internal data dictionary into a clean,
    export-ready dictionary.

    Important:
    - Experimental results are exported property-by-property.
    - AI predictions are exported property-by-property.
    - Young's Modulus is excluded.
    - Hardness is excluded.
    - Penetration-related properties are excluded.
    - No nested dictionary is written into a single Excel cell.
    """

    report = {}

    # ======================================================
    # AI PREDICTIONS
    # ======================================================

    prediction = data.get("prediction")

    if isinstance(prediction, dict):

        for prop in PREDICTION_PROPERTIES:

            if prop in prediction:

                report[prop] = prediction[prop]


    # ======================================================
    # EXPERIMENTAL RESULTS
    # ======================================================

    experimental = data.get("experimental")

    if isinstance(experimental, dict):

        for prop in EXPERIMENTAL_PROPERTIES:

            if prop in experimental:

                report[f"Experimental - {prop}"] = experimental[prop]


    # ======================================================
    # COMPARISON RESULTS
    # ======================================================

    comparison = data.get("comparison")

    if isinstance(comparison, dict):

        for prop, result in comparison.items():

            if not isinstance(result, dict):

                continue

            # ------------------------------------------------
            # Completely ignore invalid/old properties
            # ------------------------------------------------

            if prop in INVALID_PROPERTIES:

                continue

            # ------------------------------------------------
            # Only export current experimental properties
            # ------------------------------------------------

            if prop not in EXPERIMENTAL_PROPERTIES:

                continue

            report[f"{prop} - Experimental"] = result.get(
                "Experimental"
            )

            report[f"{prop} - Predicted"] = result.get(
                "Predicted"
            )

            report[f"{prop} - Error (%)"] = result.get(
                "Percentage Error (%)"
            )

            report[f"{prop} - Accuracy (%)"] = result.get(
                "Accuracy (%)"
            )


    # ======================================================
    # FORMULATION / CALCULATION DATA
    # ======================================================

    normal_keys = [
        "density",
        "epoxy",
        "horn",
        "cotton",
        "limestone",
    ]

    for key in normal_keys:

        value = data.get(key)

        if value is not None:

            report[key] = value


    return report


# ==========================================================
# EXPORT CSV
# ==========================================================

def export_to_csv(filename, data):

    with open(
        filename,
        "w",
        newline="",
        encoding="utf-8"
    ) as file:

        writer = csv.writer(file)

        writer.writerow(
            [
                "Engineering Property",
                "Value"
            ]
        )

        for key, value in data.items():

            writer.writerow(
                [
                    key,
                    value
                ]
            )

    return filename


# ==========================================================
# EXPORT EXCEL ENGINEERING REPORT
# ==========================================================

def export_to_excel(filename, data):

    workbook = Workbook()

    sheet = workbook.active

    sheet.title = "HCPS Engineering Report"


    # ======================================================
    # TITLE
    # ======================================================

    sheet["A1"] = (
        "HYBRID COMPOSITE PREDICTION SYSTEM (HCPS)"
    )

    sheet["A1"].font = Font(
        bold=True,
        size=14
    )

    sheet["A1"].alignment = Alignment(
        horizontal="center"
    )

    sheet.merge_cells("A1:B1")


    # ======================================================
    # HEADERS
    # ======================================================

    sheet["A3"] = "Engineering Property"

    sheet["B3"] = "Value"

    sheet["A3"].font = Font(
        bold=True
    )

    sheet["B3"].font = Font(
        bold=True
    )


    # ======================================================
    # DATA
    # ======================================================

    row = 4

    for key, value in data.items():

        sheet.cell(
            row=row,
            column=1,
            value=str(key)
        )

        # --------------------------------------------------
        # Preserve numeric values as numbers
        # rather than converting everything to strings.
        # --------------------------------------------------

        if isinstance(value, (int, float)):

            sheet.cell(
                row=row,
                column=2,
                value=value
            )

        else:

            sheet.cell(
                row=row,
                column=2,
                value=str(value)
            )

        row += 1


    # ======================================================
    # FORMATTING
    # ======================================================

    sheet.column_dimensions["A"].width = 45

    sheet.column_dimensions["B"].width = 25


    # ======================================================
    # SAVE WORKBOOK
    # ======================================================

    workbook.save(filename)

    return filename


# ==========================================================
# EXPORT TEXT REPORT
# ==========================================================

def export_to_text(filename, data):

    with open(
        filename,
        "w",
        encoding="utf-8"
    ) as file:

        file.write(
            "HYBRID COMPOSITE PREDICTION SYSTEM (HCPS)\n"
        )

        file.write(
            "=========================================\n\n"
        )

        for key, value in data.items():

            file.write(
                f"{key}: {value}\n"
            )

    return filename


# ==========================================================
# MAIN REPORT CREATOR
# ==========================================================

def create_report(results):

    clean_data = flatten_results(results)


    # ======================================================
    # CHECK FOR DATA
    # ======================================================

    if len(clean_data) == 0:

        raise Exception(
            "No HCPS data available for export."
        )


    # ======================================================
    # USER SELECTS SAVE LOCATION
    # ======================================================

    folder = filedialog.askdirectory(
        title="Select Folder To Save HCPS Report"
    )


    if not folder:

        raise Exception(
            "Report export cancelled."
        )


    # ======================================================
    # TIMESTAMP
    # ======================================================

    timestamp = datetime.now().strftime(
        "%Y%m%d_%H%M%S"
    )


    # ======================================================
    # FILE PATHS
    # ======================================================

    excel_file = os.path.join(
        folder,
        f"HCPS_Engineering_Report_{timestamp}.xlsx"
    )


    csv_file = os.path.join(
        folder,
        f"HCPS_Data_Report_{timestamp}.csv"
    )


    text_file = os.path.join(
        folder,
        f"HCPS_Summary_Report_{timestamp}.txt"
    )


    # ======================================================
    # EXPORT EXCEL
    # ======================================================

    export_to_excel(
        excel_file,
        clean_data
    )


    # ======================================================
    # EXPORT CSV
    # ======================================================

    export_to_csv(
        csv_file,
        clean_data
    )


    # ======================================================
    # EXPORT TEXT
    # ======================================================

    export_to_text(
        text_file,
        clean_data
    )


    # ======================================================
    # RETURN GENERATED FILES
    # ======================================================

    return (
        excel_file,
        csv_file,
        text_file
    )
